import os
import re
import datetime
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches
import pygments
from pygments.lexers import JavaLexer
from pygments.formatters import ImageFormatter

EXCEL_PATH = 'docs/SEP490_G81_Report5.1_Unit Test.xlsx'
DOCX_PATH = 'docs/SEP490_G81_Report5.1_Unit_Test_Evidence.docx'
SCRATCH_IMG_DIR = 'scratch/quote_subscription_test_images'
os.makedirs(SCRATCH_IMG_DIR, exist_ok=True)

METHOD_DEFS = [
    {
        "no": 72,
        "module": "Quote",
        "method": "getRandomQuote",
        "class_name": "GetRandomQuoteTest",
        "file_path": "core/application/src/test/java/org/naho/quote/usecase/GetRandomQuoteTest.java",
        "requirement": "Lấy ngẫu nhiên một câu châm ngôn từ cơ sở dữ liệu.",
        "description": "Lấy ngẫu nhiên một câu châm ngôn tiếng Nhật kèm phiên âm và bản dịch.",
        "precondition_summary": "1. Database có câu châm ngôn ngẫu nhiên\n2. Database không có câu châm ngôn nào",
        "cases": [
            {
                "utcid": "UTCID01_GetRandomQuote_Found_Success",
                "type": "N",
                "cond": "Database có câu châm ngôn ngẫu nhiên",
                "userData": "None",
                "mockData": "quotePort.findRandomQuote() -> Optional.of(Quote)",
                "confirm": "Lấy câu châm ngôn ngẫu nhiên thành công",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_GetRandomQuote_NotFound",
                "type": "A",
                "cond": "Database không tìm thấy câu châm ngôn nào",
                "userData": "None",
                "mockData": "quotePort.findRandomQuote() -> Optional.empty()",
                "confirm": "Ném ApplicationException mã QUOTE_NOT_FOUND (404, message: quote.not-found)",
                "checks": {"cond": 1, "userData": 0, "mockData": 1, "confirm": 1}
            }
        ],
        "conditions": [
            "Database có câu châm ngôn ngẫu nhiên",
            "Database không tìm thấy câu châm ngôn nào"
        ],
        "userDataList": [
            "None"
        ],
        "mockDataList": [
            "quotePort.findRandomQuote() -> Optional.of(Quote)",
            "quotePort.findRandomQuote() -> Optional.empty()"
        ],
        "confirmList": [
            "Lấy câu châm ngôn ngẫu nhiên thành công",
            "Ném ApplicationException mã QUOTE_NOT_FOUND (404, message: quote.not-found)"
        ]
    },
    {
        "no": 73,
        "module": "Quote",
        "method": "importQuote",
        "class_name": "ImportQuoteTest",
        "file_path": "core/application/src/test/java/org/naho/quote/usecase/ImportQuoteTest.java",
        "requirement": "Nhập danh sách câu châm ngôn từ file Excel vào hệ thống.",
        "description": "Parse và lưu danh sách câu châm ngôn tiếng Nhật từ file Excel.",
        "precondition_summary": "1. File Excel hợp lệ chứa danh sách câu châm ngôn\n2. File Excel rỗng không có câu châm ngôn nào được parse",
        "cases": [
            {
                "utcid": "UTCID01_ImportQuote_Success",
                "type": "N",
                "cond": "File Excel hợp lệ chứa danh sách câu châm ngôn",
                "userData": "InputStream inputStream = mock(InputStream.class)",
                "mockData": "quoteExcelParserPort.parseQuoteExcel(inputStream) -> List.of(quote1, quote2), quotePort.saveAll(quotes)",
                "confirm": "Nhập danh sách câu châm ngôn từ file Excel thành công",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_ImportQuote_EmptyList",
                "type": "A",
                "cond": "File Excel rỗng không có câu châm ngôn nào được parse",
                "userData": "InputStream inputStream = mock(InputStream.class)",
                "mockData": "quoteExcelParserPort.parseQuoteExcel(inputStream) -> Collections.emptyList()",
                "confirm": "Ném ApplicationException mã QUOTE_IMPORT_EMPTY (400, message: quote.import.empty)",
                "checks": {"cond": 1, "userData": 0, "mockData": 1, "confirm": 1}
            }
        ],
        "conditions": [
            "File Excel hợp lệ chứa danh sách câu châm ngôn",
            "File Excel rỗng không có câu châm ngôn nào được parse"
        ],
        "userDataList": [
            "InputStream inputStream = mock(InputStream.class)"
        ],
        "mockDataList": [
            "quoteExcelParserPort.parseQuoteExcel(inputStream) -> List.of(quotes), quotePort.saveAll(quotes)",
            "quoteExcelParserPort.parseQuoteExcel(inputStream) -> Collections.emptyList()"
        ],
        "confirmList": [
            "Nhập danh sách câu châm ngôn từ file Excel thành công",
            "Ném ApplicationException mã QUOTE_IMPORT_EMPTY (400, message: quote.import.empty)"
        ]
    },
    {
        "no": 74,
        "module": "Subscription",
        "method": "findTodayUserDailyAiUsage",
        "class_name": "FindTodayUserDailyAiUsageTest",
        "file_path": "core/application/src/test/java/org/naho/subscription/usecase/FindTodayUserDailyAiUsageTest.java",
        "requirement": "Lấy lượt sử dụng tính năng AI trong ngày của người dùng.",
        "description": "Lấy hoặc khởi tạo bản ghi sử dụng AI của người dùng trong ngày hiện tại.",
        "precondition_summary": "1. userId hợp lệ, tìm hoặc khởi tạo usage của ngày hôm nay thành công",
        "cases": [
            {
                "utcid": "UTCID01_FindTodayUserDailyAiUsage_Success",
                "type": "N",
                "cond": "userId hợp lệ, tìm hoặc khởi tạo usage của ngày hôm nay thành công",
                "userData": "userId = 1L",
                "mockData": "userDailyAiUsageRepositoryPort.findByUserIdAndUsageDateCreateIfNotExists(...) -> usage",
                "confirm": "Lấy thông tin lượt sử dụng AI trong ngày của người dùng thành công",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            }
        ],
        "conditions": [
            "userId hợp lệ, tìm hoặc khởi tạo usage của ngày hôm nay thành công"
        ],
        "userDataList": [
            "userId = 1L"
        ],
        "mockDataList": [
            "userDailyAiUsageRepositoryPort.findByUserIdAndUsageDateCreateIfNotExists(1L, today) -> usage, mapper.domainToResult(usage) -> result"
        ],
        "confirmList": [
            "Lấy thông tin lượt sử dụng AI trong ngày của người dùng thành công"
        ]
    },
    {
        "no": 75,
        "module": "Subscription",
        "method": "getUserActiveSubscriptionPlan",
        "class_name": "GetUserActiveSubscriptionPlanTest",
        "file_path": "core/application/src/test/java/org/naho/subscription/usecase/GetUserActiveSubscriptionPlanTest.java",
        "requirement": "Lấy thông tin gói cước dịch vụ đang hoạt động của người dùng.",
        "description": "Lấy gói cước dịch vụ đang active của người dùng hoặc fallback về gói FREE mặc định.",
        "precondition_summary": "1. Người dùng có gói dịch vụ trả phí đang active\n2. Người dùng không có gói trả phí active -> fallback lấy gói FREE\n3. Không có gói active và không tìm thấy gói FREE trong database",
        "cases": [
            {
                "utcid": "UTCID01_GetUserActiveSubscriptionPlan_HasActivePlan_Success",
                "type": "N",
                "cond": "Người dùng có gói dịch vụ trả phí đang active",
                "userData": "userId = 1L",
                "mockData": "subscriptionPlanRepositoryPort.findCurrentSubscriptionPlanByUserIdAndStatus(...) -> Optional.of(plan)",
                "confirm": "Lấy thông tin gói dịch vụ đang hoạt động thành công",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_GetUserActiveSubscriptionPlan_NoActivePlan_FallbackFreePlan_Success",
                "type": "N",
                "cond": "Người dùng không có gói trả phí active -> fallback lấy gói FREE",
                "userData": "userId = 1L",
                "mockData": "findCurrentSubscriptionPlanByUserIdAndStatus(...) -> Optional.empty(), findByCode(FREE) -> Optional.of(freePlan)",
                "confirm": "Fallback trả về gói cước FREE mặc định thành công",
                "checks": {"cond": 1, "userData": 0, "mockData": 1, "confirm": 1}
            },
            {
                "utcid": "UTCID03_GetUserActiveSubscriptionPlan_FreePlanNotFound",
                "type": "A",
                "cond": "Không có gói active và không tìm thấy gói FREE trong database",
                "userData": "userId = 1L",
                "mockData": "findCurrentSubscriptionPlanByUserIdAndStatus(...) -> Optional.empty(), findByCode(FREE) -> Optional.empty()",
                "confirm": "Ném ApplicationException mã PLAN_NOT_FOUND (404, message: subscription.plan.not-found)",
                "checks": {"cond": 2, "userData": 0, "mockData": 2, "confirm": 2}
            }
        ],
        "conditions": [
            "Người dùng có gói dịch vụ trả phí đang active",
            "Người dùng không có gói trả phí active -> fallback lấy gói FREE",
            "Không có gói active và không tìm thấy gói FREE trong database"
        ],
        "userDataList": [
            "userId = 1L"
        ],
        "mockDataList": [
            "subscriptionPlanRepositoryPort.findCurrentSubscriptionPlanByUserIdAndStatus(...) -> Optional.of(activePlan)",
            "findCurrentSubscriptionPlanByUserIdAndStatus(...) -> Optional.empty(), findByCode(PlanCode.FREE) -> Optional.of(freePlan)",
            "findCurrentSubscriptionPlanByUserIdAndStatus(...) -> Optional.empty(), findByCode(PlanCode.FREE) -> Optional.empty()"
        ],
        "confirmList": [
            "Lấy thông tin gói dịch vụ đang hoạt động thành công",
            "Fallback trả về gói cước FREE mặc định thành công",
            "Ném ApplicationException mã PLAN_NOT_FOUND (404, message: subscription.plan.not-found)"
        ]
    },
    {
        "no": 76,
        "module": "Subscription",
        "method": "getUserActiveSubscription",
        "class_name": "GetUserActiveSubscriptionTest",
        "file_path": "core/application/src/test/java/org/naho/subscription/usecase/GetUserActiveSubscriptionTest.java",
        "requirement": "Lấy chi tiết bản ghi đăng ký gói cước dịch vụ của người dùng.",
        "description": "Lấy thông tin đăng ký gói cước active hoặc trả về thông tin mặc định nếu dùng gói FREE.",
        "precondition_summary": "1. Người dùng có bản ghi UserSubscription đang active\n2. Người dùng dùng gói FREE mặc định (chưa có bản ghi UserSubscription)",
        "cases": [
            {
                "utcid": "UTCID01_GetUserActiveSubscription_HasUserSubscription_Success",
                "type": "N",
                "cond": "Người dùng có bản ghi UserSubscription đang active",
                "userData": "userId = 1L",
                "mockData": "getUserActiveSubscriptionPlan(1L) -> planResult, userSubscriptionRepositoryPort.findActiveByUserId(...) -> Optional.of(userSub)",
                "confirm": "Lấy thông tin gói đăng ký kèm chi tiết đăng ký thành công",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_GetUserActiveSubscription_DefaultFreeNoSubscription_Success",
                "type": "N",
                "cond": "Người dùng dùng gói FREE mặc định (chưa có bản ghi UserSubscription)",
                "userData": "userId = 1L",
                "mockData": "getUserActiveSubscriptionPlan(1L) -> freePlanResult, userSubscriptionRepositoryPort.findActiveByUserId(...) -> Optional.empty()",
                "confirm": "Trả về UserSubscriptionResult mặc định cho gói FREE thành công",
                "checks": {"cond": 1, "userData": 0, "mockData": 1, "confirm": 1}
            }
        ],
        "conditions": [
            "Người dùng có bản ghi UserSubscription đang active",
            "Người dùng dùng gói FREE mặc định (chưa có bản ghi UserSubscription)"
        ],
        "userDataList": [
            "userId = 1L"
        ],
        "mockDataList": [
            "getUserActiveSubscriptionPlan(1L) -> planResult, userSubscriptionRepositoryPort.findActiveByUserId(...) -> Optional.of(userSub)",
            "getUserActiveSubscriptionPlan(1L) -> freePlanResult, userSubscriptionRepositoryPort.findActiveByUserId(...) -> Optional.empty()"
        ],
        "confirmList": [
            "Lấy thông tin gói đăng ký kèm chi tiết đăng ký thành công",
            "Trả về UserSubscriptionResult mặc định cho gói FREE thành công"
        ]
    },
    {
        "no": 77,
        "module": "Subscription",
        "method": "listActivePlans",
        "class_name": "ListActivePlansTest",
        "file_path": "core/application/src/test/java/org/naho/subscription/usecase/ListActivePlansTest.java",
        "requirement": "Lấy danh sách tất cả các gói dịch vụ đang hoạt động.",
        "description": "Lấy danh sách các gói cước dịch vụ có trạng thái ACTIVE trong hệ thống.",
        "precondition_summary": "1. Hệ thống có danh sách các gói dịch vụ active\n2. Hệ thống không có gói dịch vụ nào active",
        "cases": [
            {
                "utcid": "UTCID01_ListActivePlans_HasData_Success",
                "type": "N",
                "cond": "Hệ thống có danh sách các gói dịch vụ active",
                "userData": "None",
                "mockData": "planRepositoryPort.findAllActive() -> List.of(plan1, plan2)",
                "confirm": "Lấy danh sách các gói dịch vụ đang hoạt động thành công",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_ListActivePlans_EmptyList_Success",
                "type": "N",
                "cond": "Hệ thống không có gói dịch vụ nào active",
                "userData": "None",
                "mockData": "planRepositoryPort.findAllActive() -> Collections.emptyList()",
                "confirm": "Trả về danh sách rỗng khi không có gói cước nào hoạt động",
                "checks": {"cond": 1, "userData": 0, "mockData": 1, "confirm": 1}
            }
        ],
        "conditions": [
            "Hệ thống có danh sách các gói dịch vụ active",
            "Hệ thống không có gói dịch vụ nào active"
        ],
        "userDataList": [
            "None"
        ],
        "mockDataList": [
            "planRepositoryPort.findAllActive() -> List.of(plan1, plan2)",
            "planRepositoryPort.findAllActive() -> Collections.emptyList()"
        ],
        "confirmList": [
            "Lấy danh sách các gói dịch vụ đang hoạt động thành công",
            "Trả về danh sách rỗng khi không có gói cước nào hoạt động"
        ]
    },
    {
        "no": 78,
        "module": "Subscription",
        "method": "updateSubscriptionPlan",
        "class_name": "UpdateSubscriptionPlanTest",
        "file_path": "core/application/src/test/java/org/naho/subscription/usecase/UpdateSubscriptionPlanTest.java",
        "requirement": "Cập nhật thông tin chi tiết của gói cước dịch vụ.",
        "description": "ADMIN cập nhật thông tin mô tả, giá cả và các thông số hạn mức của gói cước.",
        "precondition_summary": "1. Người thực hiện không có quyền ADMIN\n2. Người thực hiện là ADMIN nhưng không tìm thấy gói cước theo ID\n3. ADMIN cập nhật thành công thông tin gói cước bao gồm giá mới và các giới hạn\n4. Xác thực quyền ADMIN qua đối tượng Role và cập nhật gói cước thành công",
        "cases": [
            {
                "utcid": "UTCID01_UpdateSubscriptionPlan_AccessDenied",
                "type": "A",
                "cond": "Người thực hiện không có quyền ADMIN",
                "userData": "UpdateSubscriptionPlanCommand(planId=1L, adminUserId=10L, ...)",
                "mockData": "roleRepositoryPort.findRoleNamesByUserId(10L) -> List.of('LEARNER'), findByUserId(10L) -> Role.LEARNER",
                "confirm": "Ném ApplicationException mã USER_ACCESS_DENIED (403, message: user.access.denied)",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_UpdateSubscriptionPlan_PlanNotFound",
                "type": "A",
                "cond": "Người thực hiện là ADMIN nhưng không tìm thấy gói cước theo ID",
                "userData": "UpdateSubscriptionPlanCommand(planId=999L, adminUserId=1L, ...)",
                "mockData": "roleRepositoryPort.findRoleNamesByUserId(1L) -> List.of('ADMIN'), planRepositoryPort.findById(999L) -> Optional.empty()",
                "confirm": "Ném ApplicationException mã PLAN_NOT_FOUND (404, message: subscription.plan.not-found)",
                "checks": {"cond": 1, "userData": 1, "mockData": 1, "confirm": 1}
            },
            {
                "utcid": "UTCID03_UpdateSubscriptionPlan_Success",
                "type": "N",
                "cond": "ADMIN cập nhật thành công thông tin gói cước bao gồm giá mới và các giới hạn",
                "userData": "UpdateSubscriptionPlanCommand(planId=2L, adminUserId=1L, priceAmount=250000, ...)",
                "mockData": "roleRepositoryPort.findRoleNamesByUserId(1L) -> List.of('ADMIN'), planRepositoryPort.findById(2L) -> Optional.of(plan), save(plan)",
                "confirm": "Cập nhật thông tin gói cước thành công",
                "checks": {"cond": 2, "userData": 2, "mockData": 2, "confirm": 2}
            },
            {
                "utcid": "UTCID04_UpdateSubscriptionPlan_AdminVerifiedViaRoleObject_Success",
                "type": "N",
                "cond": "Xác thực quyền ADMIN qua đối tượng Role và cập nhật gói cước thành công",
                "userData": "UpdateSubscriptionPlanCommand(planId=2L, adminUserId=1L, priceAmount=null, ...)",
                "mockData": "findRoleNamesByUserId(1L) -> List.of(), findByUserId(1L) -> Role.ADMIN, planRepositoryPort.findById(2L) -> Optional.of(plan), save(plan)",
                "confirm": "Xác thực quyền ADMIN qua đối tượng Role và cập nhật gói cước thành công",
                "checks": {"cond": 3, "userData": 3, "mockData": 3, "confirm": 2}
            }
        ],
        "conditions": [
            "Người thực hiện không có quyền ADMIN",
            "Người thực hiện là ADMIN nhưng không tìm thấy gói cước theo ID",
            "ADMIN cập nhật thành công thông tin gói cước bao gồm giá mới và các giới hạn",
            "Xác thực quyền ADMIN qua đối tượng Role và cập nhật gói cước thành công"
        ],
        "userDataList": [
            "UpdateSubscriptionPlanCommand(planId=1L, adminUserId=10L, description='Description', tier=BASIC, price=199000, ...)",
            "UpdateSubscriptionPlanCommand(planId=999L, adminUserId=1L, description='Description', tier=BASIC, price=199000, ...)",
            "UpdateSubscriptionPlanCommand(planId=2L, adminUserId=1L, description='Updated Description', tier=BASIC, price=250000, ...)",
            "UpdateSubscriptionPlanCommand(planId=2L, adminUserId=1L, description='Updated Description', tier=BASIC, price=null, ...)"
        ],
        "mockDataList": [
            "roleRepositoryPort.findRoleNamesByUserId(10L) -> List.of('LEARNER'), roleRepositoryPort.findByUserId(10L) -> Role.LEARNER",
            "roleRepositoryPort.findRoleNamesByUserId(1L) -> List.of('ADMIN'), planRepositoryPort.findById(999L) -> Optional.empty()",
            "roleRepositoryPort.findRoleNamesByUserId(1L) -> List.of('ADMIN'), planRepositoryPort.findById(2L) -> Optional.of(plan), save(plan) -> updated",
            "roleRepositoryPort.findRoleNamesByUserId(1L) -> List.of(), roleRepositoryPort.findByUserId(1L) -> Role.ADMIN, planRepositoryPort.findById(2L) -> plan, save(plan)"
        ],
        "confirmList": [
            "Ném ApplicationException mã USER_ACCESS_DENIED (403, message: user.access.denied)",
            "Ném ApplicationException mã PLAN_NOT_FOUND (404, message: subscription.plan.not-found)",
            "Cập nhật thông tin gói cước thành công",
            "Xác thực quyền ADMIN qua đối tượng Role và cập nhật gói cước thành công"
        ]
    },
    {
        "no": 79,
        "module": "Subscription",
        "method": "mapToPlanResult",
        "class_name": "SubscriptionPlanResultMapperTest",
        "file_path": "core/application/src/test/java/org/naho/subscription/mapper/SubscriptionPlanResultMapperTest.java",
        "requirement": "Chuyển đổi SubscriptionPlan domain sang SubscriptionPlanResult.",
        "description": "Mapper chuyển đổi đối tượng domain SubscriptionPlan sang DTO kết quả SubscriptionPlanResult.",
        "precondition_summary": "1. Domain SubscriptionPlan hợp lệ khác null\n2. Domain SubscriptionPlan là null",
        "cases": [
            {
                "utcid": "UTCID01_MapToPlanResult_Success",
                "type": "N",
                "cond": "Domain SubscriptionPlan hợp lệ khác null",
                "userData": "SubscriptionPlan(id=1L, code=BASIC, price=199000 VND, durationDays=30, ...)",
                "mockData": "None",
                "confirm": "Chuyển đổi SubscriptionPlan domain sang SubscriptionPlanResult thành công",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_MapToPlanResult_Null",
                "type": "N",
                "cond": "Domain SubscriptionPlan là null",
                "userData": "plan = null",
                "mockData": "None",
                "confirm": "Chuyển đổi trả về null khi domain plan là null",
                "checks": {"cond": 1, "userData": 1, "mockData": 0, "confirm": 1}
            }
        ],
        "conditions": [
            "Domain SubscriptionPlan hợp lệ khác null",
            "Domain SubscriptionPlan là null"
        ],
        "userDataList": [
            "SubscriptionPlan(id=1L, code=BASIC, description='Basic Plan', tier=BASIC, price=199000 VND, durationDays=30, ...)",
            "plan = null"
        ],
        "mockDataList": [
            "None"
        ],
        "confirmList": [
            "Chuyển đổi SubscriptionPlan domain sang SubscriptionPlanResult thành công",
            "Chuyển đổi trả về null khi domain plan là null"
        ]
    },
    {
        "no": 80,
        "module": "Subscription",
        "method": "domainToResultUsage",
        "class_name": "UserDailyAiUsageResultMapperTest",
        "file_path": "core/application/src/test/java/org/naho/subscription/mapper/UserDailyAiUsageResultMapperTest.java",
        "requirement": "Chuyển đổi UserDailyAiUsage domain sang UserDailyAiUsageResult.",
        "description": "Mapper chuyển đổi đối tượng domain UserDailyAiUsage sang DTO kết quả UserDailyAiUsageResult.",
        "precondition_summary": "1. Domain UserDailyAiUsage hợp lệ khác null\n2. Domain UserDailyAiUsage là null",
        "cases": [
            {
                "utcid": "UTCID01_DomainToResult_Success",
                "type": "N",
                "cond": "Domain UserDailyAiUsage hợp lệ khác null",
                "userData": "UserDailyAiUsage(id=10L, userId=1L, usageDate=2026-08-19, speakingEvaluationCount=5, aiSessionStartCount=2)",
                "mockData": "None",
                "confirm": "Chuyển đổi UserDailyAiUsage domain sang UserDailyAiUsageResult thành công",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_DomainToResult_Null",
                "type": "N",
                "cond": "Domain UserDailyAiUsage là null",
                "userData": "usage = null",
                "mockData": "None",
                "confirm": "Chuyển đổi trả về null khi domain usage là null",
                "checks": {"cond": 1, "userData": 1, "mockData": 0, "confirm": 1}
            }
        ],
        "conditions": [
            "Domain UserDailyAiUsage hợp lệ khác null",
            "Domain UserDailyAiUsage là null"
        ],
        "userDataList": [
            "UserDailyAiUsage(id=10L, userId=1L, usageDate=2026-08-19, speakingEvaluationCount=5, aiSessionStartCount=2)",
            "usage = null"
        ],
        "mockDataList": [
            "None"
        ],
        "confirmList": [
            "Chuyển đổi UserDailyAiUsage domain sang UserDailyAiUsageResult thành công",
            "Chuyển đổi trả về null khi domain usage là null"
        ]
    },
    {
        "no": 81,
        "module": "Subscription",
        "method": "mapToUserSubscriptionResult",
        "class_name": "UserSubscriptionResultMapperTest",
        "file_path": "core/application/src/test/java/org/naho/subscription/mapper/UserSubscriptionResultMapperTest.java",
        "requirement": "Chuyển đổi UserSubscription domain sang UserSubscriptionResult.",
        "description": "Mapper chuyển đổi đối tượng domain UserSubscription và SubscriptionPlanResult sang UserSubscriptionResult.",
        "precondition_summary": "1. Domain UserSubscription hợp lệ khác null\n2. Domain UserSubscription là null",
        "cases": [
            {
                "utcid": "UTCID01_MapToUserSubscriptionResult_Success",
                "type": "N",
                "cond": "Domain UserSubscription hợp lệ khác null",
                "userData": "UserSubscription(id=100L, userId=1L, planId=2L, paymentOrderId=50L, status=ACTIVE, ...)",
                "mockData": "planResult = mock(SubscriptionPlanResult.class)",
                "confirm": "Chuyển đổi UserSubscription domain sang UserSubscriptionResult thành công",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_MapToUserSubscriptionResult_Null",
                "type": "N",
                "cond": "Domain UserSubscription là null",
                "userData": "subscription = null",
                "mockData": "None",
                "confirm": "Chuyển đổi trả về null khi domain subscription là null",
                "checks": {"cond": 1, "userData": 1, "mockData": 0, "confirm": 1}
            }
        ],
        "conditions": [
            "Domain UserSubscription hợp lệ khác null",
            "Domain UserSubscription là null"
        ],
        "userDataList": [
            "UserSubscription(id=100L, userId=1L, planId=2L, paymentOrderId=50L, status=ACTIVE, ...)",
            "subscription = null"
        ],
        "mockDataList": [
            "planResult = mock(SubscriptionPlanResult.class)",
            "None"
        ],
        "confirmList": [
            "Chuyển đổi UserSubscription domain sang UserSubscriptionResult thành công",
            "Chuyển đổi trả về null khi domain subscription là null"
        ]
    }
]

def to_roman(n):
    val = [1000, 900, 500, 400, 100, 90, 50, 40, 10, 9, 5, 4, 1]
    syb = ["m", "cm", "d", "cd", "c", "xc", "l", "xl", "x", "ix", "v", "iv", "i"]
    res = ""
    for i in range(len(val)):
        while n >= val[i]:
            res += syb[i]
            n -= val[i]
    return res

def create_excel_sheet(wb, mdef):
    sample = wb['sample']
    sheet_name = mdef['method']
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.copy_worksheet(sample)
    ws.title = sheet_name

    # Set headers
    ws['C1'] = mdef['module']
    ws['L1'] = mdef['method']
    ws['C2'] = 'TrucNV'
    ws['L2'] = 'TrucNV'
    ws['C3'] = mdef['requirement']

    num_cases = len(mdef['cases'])
    start_col = 6  # Col F
    end_col = start_col + num_cases - 1
    start_col_letter = get_column_letter(start_col)
    end_col_letter = get_column_letter(end_col)

    # Clear old data from row 7 downwards
    for r in range(7, 50):
        for c in range(1, 30):
            ws.cell(r, c).value = None

    # Row 7: Test Case IDs
    for idx, cdef in enumerate(mdef['cases']):
        col = start_col + idx
        ws.cell(7, col, f"UTCID{idx+1:02d}")

    curr_row = 8
    # 1. Condition
    ws.cell(curr_row, 1, 'Condition')
    ws.cell(curr_row, 2, 'Precondition')
    curr_row += 1

    for idx, cond_text in enumerate(mdef['conditions']):
        ws.cell(curr_row, 2, cond_text)
        for c_idx, cdef in enumerate(mdef['cases']):
            if cdef['checks']['cond'] == idx:
                ws.cell(curr_row, start_col + c_idx, 'O')
        curr_row += 1

    # Blank row
    curr_row += 1

    # 2. userData
    ws.cell(curr_row, 1, 'userData')
    ws.cell(curr_row, 2, 'Command Data')
    curr_row += 1

    for idx, udata_text in enumerate(mdef['userDataList']):
        ws.cell(curr_row, 2, udata_text)
        for c_idx, cdef in enumerate(mdef['cases']):
            if cdef['checks']['userData'] == idx:
                ws.cell(curr_row, start_col + c_idx, 'O')
        curr_row += 1

    # Blank row
    curr_row += 1

    # 3. mockData
    ws.cell(curr_row, 1, 'mockData')
    ws.cell(curr_row, 2, 'Repository / Service')
    curr_row += 1

    for idx, mdata_text in enumerate(mdef['mockDataList']):
        ws.cell(curr_row, 2, mdata_text)
        for c_idx, cdef in enumerate(mdef['cases']):
            if cdef['checks']['mockData'] == idx:
                ws.cell(curr_row, start_col + c_idx, 'O')
        curr_row += 1

    # Blank row
    curr_row += 1

    # 4. Confirm
    ws.cell(curr_row, 1, 'Confirm')
    ws.cell(curr_row, 2, 'Expected Result')
    curr_row += 1

    for idx, conf_text in enumerate(mdef['confirmList']):
        ws.cell(curr_row, 2, conf_text)
        for c_idx, cdef in enumerate(mdef['cases']):
            if cdef['checks']['confirm'] == idx:
                ws.cell(curr_row, start_col + c_idx, 'O')
        curr_row += 1

    # Blank row
    curr_row += 1

    # 5. Result
    res_start_row = curr_row
    ws.cell(curr_row, 1, 'Result')
    ws.cell(curr_row, 2, 'Type')
    for c_idx, cdef in enumerate(mdef['cases']):
        ws.cell(curr_row, start_col + c_idx, cdef['type'])
    type_row = curr_row
    curr_row += 1

    ws.cell(curr_row, 2, 'Passed/Failed')
    for c_idx, cdef in enumerate(mdef['cases']):
        ws.cell(curr_row, start_col + c_idx, 'P')
    pf_row = curr_row
    curr_row += 1

    ws.cell(curr_row, 2, 'Executed Date')
    for c_idx, cdef in enumerate(mdef['cases']):
        ws.cell(curr_row, start_col + c_idx, datetime.datetime(2026, 8, 19, 0, 0))
    curr_row += 1

    ws.cell(curr_row, 2, 'Defect ID')
    curr_row += 1

    # Formulas in Row 5
    ws['A5'] = f'=COUNTIF({start_col_letter}{pf_row}:{end_col_letter}{pf_row},"P")'
    ws['C5'] = f'=COUNTIF({start_col_letter}{pf_row}:{end_col_letter}{pf_row},"F")'
    ws['F5'] = f'=COUNTIF({start_col_letter}{pf_row}:{end_col_letter}{pf_row},"U")'
    ws['L5'] = f'=COUNTIF({start_col_letter}{type_row}:{end_col_letter}{type_row},"N")'
    ws['M5'] = f'=COUNTIF({start_col_letter}{type_row}:{end_col_letter}{type_row},"A")'
    ws['N5'] = f'=COUNTIF({start_col_letter}{type_row}:{end_col_letter}{type_row},"B")'
    ws['O5'] = f'=COUNTA({start_col_letter}7:{end_col_letter}7)'

    # Add Data Validations
    dv_o = DataValidation(type="list", formula1='"O"', allow_blank=True)
    ws.add_data_validation(dv_o)
    dv_o.add(f'{start_col_letter}8:{end_col_letter}{type_row - 2}')

    dv_type = DataValidation(type="list", formula1='"N,A,B"', allow_blank=True)
    ws.add_data_validation(dv_type)
    dv_type.add(f'{start_col_letter}{type_row}:{end_col_letter}{type_row}')

    dv_pf = DataValidation(type="list", formula1='"P,F,U"', allow_blank=True)
    ws.add_data_validation(dv_pf)
    dv_pf.add(f'{start_col_letter}{pf_row}:{end_col_letter}{pf_row}')

    print(f'Created sheet {sheet_name} with {num_cases} test cases.')

def update_statistics(wb):
    ws = wb['Statistics']
    ws['F6'] = datetime.datetime(2026, 8, 19, 0, 0)

    for idx, mdef in enumerate(METHOD_DEFS):
        r = 82 + idx
        s_name = mdef['method']
        ws.cell(r, 1, mdef['no'])
        ws.cell(r, 2, f'=HYPERLINK("#{s_name}!A1", MethodList!C{r})')
        ws.cell(r, 3, f'={s_name}!A5')
        ws.cell(r, 4, f'={s_name}!C5')
        ws.cell(r, 5, f'={s_name}!F5')
        ws.cell(r, 6, f'={s_name}!L5')
        ws.cell(r, 7, f'={s_name}!M5')
        ws.cell(r, 8, f'={s_name}!N5')
        ws.cell(r, 9, f'=SUM(C{r}:H{r})')
    print('Updated Statistics sheet.')

def update_method_list(wb):
    ws = wb['MethodList']
    for idx, mdef in enumerate(METHOD_DEFS):
        r = 82 + idx
        s_name = mdef['method']
        ws.cell(r, 1, mdef['no'])
        ws.cell(r, 2, mdef['module'])
        ws.cell(r, 3, mdef['method'])
        ws.cell(r, 4, f'=HYPERLINK("#{s_name}!A1", "{s_name}")')
        ws.cell(r, 5, mdef['description'])
        ws.cell(r, 6, mdef['precondition_summary'])
    print('Updated MethodList sheet.')

def update_cover(wb):
    ws = wb['Cover']
    ws['F5'] = datetime.datetime(2026, 8, 19, 0, 0)
    ws['F6'] = datetime.datetime(2026, 8, 19, 0, 0)

    # Add change record for Quote & Subscription
    r = 82
    ws.cell(r, 1, datetime.datetime(2026, 8, 19, 0, 0))
    ws.cell(r, 2, '1.0')
    ws.cell(r, 3, 'Quote, Subscription')
    ws.cell(r, 4, 'A')
    ws.cell(r, 5, 'Add Unit Test for Quote and Subscription modules')
    ws.cell(r, 6, 'Quote, Subscription')
    print('Updated Cover sheet.')

def generate_code_images():
    formatter = ImageFormatter(
        style='perldoc',
        font_name='Menlo',
        font_size=15,
        line_pad=6
    )

    img_map = {}
    for mdef in METHOD_DEFS:
        fpath = mdef['file_path']
        with open(fpath, 'r', encoding='utf-8') as f:
            content = f.read()

        pattern = re.compile(r'(\s+@Test[\s\S]*?\n    \})', re.MULTILINE)
        matches = pattern.findall(content)
        for m in matches:
            name_match = re.search(r'void\s+([A-Za-z0-9_]+)\s*\(', m)
            if name_match:
                mname = name_match.group(1)
                lines = m.split('\n')
                cleaned_lines = []
                for line in lines:
                    if line.startswith('    '):
                        cleaned_lines.append(line[4:])
                    else:
                        cleaned_lines.append(line)
                cleaned_code = '\n'.join(cleaned_lines).strip()
                img_data = pygments.highlight(cleaned_code, JavaLexer(), formatter)
                out_path = os.path.join(SCRATCH_IMG_DIR, f"{mname}.png")
                with open(out_path, 'wb') as img_f:
                    img_f.write(img_data)
                img_map[mname] = out_path
                print(f'Generated image for {mname} -> {out_path}')
    return img_map

def update_docx(img_map):
    doc = docx.Document(DOCX_PATH)

    for mdef in METHOD_DEFS:
        h2_text = f"{mdef['no']}. {mdef['class_name']}"
        doc.add_paragraph(h2_text, style='Heading 2')

        for idx, cdef in enumerate(mdef['cases']):
            roman = to_roman(idx + 1)
            utcid = cdef['utcid']
            h3_text = f"{roman}. {utcid}"
            doc.add_paragraph(h3_text, style='Heading 3')

            img_path = img_map.get(utcid)
            if img_path and os.path.exists(img_path):
                p_img = doc.add_paragraph()
                r_img = p_img.add_run()
                r_img.add_picture(img_path, width=Inches(6.2))
            else:
                print(f'WARNING: Image for {utcid} not found!')

    doc.save(DOCX_PATH)
    print('Updated docx successfully.')

def main():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    for mdef in METHOD_DEFS:
        create_excel_sheet(wb, mdef)
    update_statistics(wb)
    update_method_list(wb)
    update_cover(wb)
    wb.save(EXCEL_PATH)
    print('Saved Excel file.')

    img_map = generate_code_images()
    update_docx(img_map)

if __name__ == '__main__':
    main()
