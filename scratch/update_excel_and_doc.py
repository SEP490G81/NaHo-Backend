import os
import re
import datetime
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches
import pygments
from pygments.lexers import JavaLexer
from pygments.formatters import ImageFormatter

EXCEL_PATH = 'docs/SEP490_G81_Report5.1_Unit Test.xlsx'
DOCX_PATH = 'docs/SEP490_G81_Report5.1_Unit_Test_Evidence.docx'
TEST_DIR = 'core/application/src/test/java/org/naho/persona/usecase'
SCRATCH_IMG_DIR = 'scratch/persona_test_images'
os.makedirs(SCRATCH_IMG_DIR, exist_ok=True)

METHOD_DEFS = [
    {
        "no": 67,
        "module": "Persona",
        "method": "createPersona",
        "class_name": "CreatePersonaTest",
        "file_name": "CreatePersonaTest.java",
        "requirement": "Tạo mới nhân vật và cấu hình phong cách hội thoại.",
        "description": "Tạo mới một nhân vật (Persona) kèm hoặc không kèm tạo mới phong cách hội thoại (ConversationStyle).",
        "precondition_summary": "1. Command có suggestedConversationStyleId có sẵn\n2. Command có conversationStyleCommand để tạo phong cách mới\n3. Command không có styleId và không có styleCommand",
        "cases": [
            {
                "utcid": "UTCID01_CreatePersona_WithExistingStyleId_Success",
                "type": "N",
                "cond": "Command hợp lệ, có suggestedConversationStyleId có sẵn, conversationStyleCommand = null",
                "userData": "name = \"Sensei Tanaka\", prompt = \"You are a friendly Japanese teacher\", avatarFileId = 10L, suggestedConversationStyleId = 1L",
                "mockData": "personaRepositoryPort.save(Persona) -> Persona(id=1L)",
                "confirm": "Tạo mới nhân vật thành công với ID phong cách hội thoại có sẵn",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_CreatePersona_WithNewConversationStyleCommand_Success",
                "type": "N",
                "cond": "Command hợp lệ, có conversationStyleCommand để tạo phong cách mới",
                "userData": "name = \"Anime Hero\", prompt = \"Energetic anime hero\", avatarFileId = 20L, conversationStyleCommand = CreateConversationStyleCommand(...)",
                "mockData": "conversationStyleRepositoryPort.save(ConversationStyle) -> id=2L, personaRepositoryPort.save(Persona) -> id=2L",
                "confirm": "Tạo mới nhân vật kèm phong cách hội thoại mới thành công",
                "checks": {"cond": 1, "userData": 1, "mockData": 1, "confirm": 1}
            },
            {
                "utcid": "UTCID03_CreatePersona_WithoutStyle_Success",
                "type": "N",
                "cond": "Command hợp lệ, không có styleId và không có conversationStyleCommand",
                "userData": "name = \"Plain Persona\", prompt = \"A plain assistant\", avatarFileId = null, suggestedConversationStyleId = null",
                "mockData": "personaRepositoryPort.save(Persona) -> Persona(id=3L)",
                "confirm": "Tạo mới nhân vật không kèm phong cách hội thoại thành công",
                "checks": {"cond": 2, "userData": 2, "mockData": 2, "confirm": 2}
            }
        ],
        "conditions": [
            "Command hợp lệ, có suggestedConversationStyleId có sẵn, conversationStyleCommand = null",
            "Command hợp lệ, có conversationStyleCommand để tạo phong cách mới",
            "Command hợp lệ, không có styleId và không có conversationStyleCommand"
        ],
        "userDataList": [
            "CreatePersonaCommand(name='Sensei Tanaka', prompt='You are...', avatarFileId=10L, suggestedConversationStyleId=1L, conversationStyleCommand=null)",
            "CreatePersonaCommand(name='Anime Hero', prompt='Energetic...', avatarFileId=20L, conversationStyleCommand=CreateConversationStyleCommand(...))",
            "CreatePersonaCommand(name='Plain Persona', prompt='A plain...', avatarFileId=null, suggestedConversationStyleId=null, conversationStyleCommand=null)"
        ],
        "mockDataList": [
            "personaRepositoryPort.save(any(Persona.class)) -> savedPersona",
            "conversationStyleRepositoryPort.save(any(ConversationStyle.class)) -> savedStyle",
            "personaRepositoryPort.save(any(Persona.class)) -> savedPersona (no style)"
        ],
        "confirmList": [
            "Tạo mới nhân vật thành công với ID phong cách hội thoại có sẵn",
            "Tạo mới nhân vật kèm phong cách hội thoại mới thành công",
            "Tạo mới nhân vật không kèm phong cách hội thoại thành công"
        ]
    },
    {
        "no": 68,
        "module": "Persona",
        "method": "getAllPersonas",
        "class_name": "GetAllPersonasTest",
        "file_name": "GetAllPersonasTest.java",
        "requirement": "Lấy toàn bộ danh sách nhân vật có trong hệ thống.",
        "description": "Lấy danh sách tất cả các nhân vật đã được tạo trong hệ thống.",
        "precondition_summary": "1. Hệ thống có danh sách nhân vật trong cơ sở dữ liệu\n2. Hệ thống chưa có nhân vật nào trong cơ sở dữ liệu",
        "cases": [
            {
                "utcid": "UTCID01_GetAllPersonas_HasData_Success",
                "type": "N",
                "cond": "Hệ thống có danh sách nhân vật trong cơ sở dữ liệu",
                "userData": "None",
                "mockData": "personaRepositoryPort.findAll() -> List.of(Persona1, Persona2)",
                "confirm": "Lấy danh sách tất cả nhân vật thành công",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_GetAllPersonas_Empty_Success",
                "type": "N",
                "cond": "Hệ thống chưa có nhân vật nào trong cơ sở dữ liệu",
                "userData": "None",
                "mockData": "personaRepositoryPort.findAll() -> Collections.emptyList()",
                "confirm": "Lấy danh sách nhân vật trả về danh sách rỗng",
                "checks": {"cond": 1, "userData": 0, "mockData": 1, "confirm": 1}
            }
        ],
        "conditions": [
            "Hệ thống có danh sách nhân vật trong cơ sở dữ liệu",
            "Hệ thống chưa có nhân vật nào trong cơ sở dữ liệu"
        ],
        "userDataList": [
            "None"
        ],
        "mockDataList": [
            "personaRepositoryPort.findAll() -> List.of(Persona1, Persona2)",
            "personaRepositoryPort.findAll() -> Collections.emptyList()"
        ],
        "confirmList": [
            "Lấy danh sách tất cả nhân vật thành công",
            "Lấy danh sách nhân vật trả về danh sách rỗng"
        ]
    },
    {
        "no": 69,
        "module": "Persona",
        "method": "getPersonaById",
        "class_name": "GetPersonaByIdTest",
        "file_name": "GetPersonaByIdTest.java",
        "requirement": "Lấy thông tin chi tiết của nhân vật theo ID.",
        "description": "Lấy chi tiết một nhân vật theo định danh ID.",
        "precondition_summary": "1. Nhân vật tồn tại trong cơ sở dữ liệu\n2. Nhân vật không tồn tại trong cơ sở dữ liệu",
        "cases": [
            {
                "utcid": "UTCID01_GetPersonaById_Found_Success",
                "type": "N",
                "cond": "Nhân vật tồn tại trong cơ sở dữ liệu",
                "userData": "id = 1L",
                "mockData": "personaRepositoryPort.findById(1L) -> Optional.of(Persona)",
                "confirm": "Lấy thông tin nhân vật theo ID thành công",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_GetPersonaById_NotFound",
                "type": "A",
                "cond": "Nhân vật không tồn tại trong cơ sở dữ liệu",
                "userData": "id = 999L",
                "mockData": "personaRepositoryPort.findById(999L) -> Optional.empty()",
                "confirm": "Không tìm thấy nhân vật, trả về Optional rỗng",
                "checks": {"cond": 1, "userData": 1, "mockData": 1, "confirm": 1}
            }
        ],
        "conditions": [
            "Nhân vật tồn tại trong cơ sở dữ liệu",
            "Nhân vật không tồn tại trong cơ sở dữ liệu"
        ],
        "userDataList": [
            "id = 1L",
            "id = 999L"
        ],
        "mockDataList": [
            "personaRepositoryPort.findById(1L) -> Optional.of(persona)",
            "personaRepositoryPort.findById(999L) -> Optional.empty()"
        ],
        "confirmList": [
            "Lấy thông tin nhân vật theo ID thành công",
            "Không tìm thấy nhân vật, trả về Optional rỗng"
        ]
    },
    {
        "no": 70,
        "module": "Persona",
        "method": "getConversationStyleByPersonaId",
        "class_name": "GetConversationStyleByPersonaIdTest",
        "file_name": "GetConversationStyleByPersonaIdTest.java",
        "requirement": "Lấy thông tin phong cách hội thoại của nhân vật theo ID.",
        "description": "Lấy phong cách hội thoại được liên kết hoặc gợi ý cho nhân vật.",
        "precondition_summary": "1. Nhân vật không tồn tại trong cơ sở dữ liệu\n2. Nhân vật tồn tại và đối tượng phong cách hội thoại liên kết trực tiếp\n3. Nhân vật tồn tại và có suggestedConversationStyleId hợp lệ\n4. Nhân vật tồn tại nhưng không có phong cách hội thoại nào được cấu hình",
        "cases": [
            {
                "utcid": "UTCID01_GetConversationStyleByPersonaId_PersonaNotFound",
                "type": "A",
                "cond": "Nhân vật không tồn tại trong cơ sở dữ liệu",
                "userData": "id = 999L",
                "mockData": "personaRepositoryPort.findById(999L) -> Optional.empty()",
                "confirm": "Không tìm thấy nhân vật, trả về Optional rỗng",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_GetConversationStyleByPersonaId_EmbeddedStylePresent_Success",
                "type": "N",
                "cond": "Nhân vật tồn tại và đối tượng phong cách hội thoại liên kết trực tiếp",
                "userData": "id = 1L",
                "mockData": "personaRepositoryPort.findById(1L) -> Optional.of(personaWithStyle)",
                "confirm": "Lấy phong cách hội thoại từ đối tượng liên kết trực tiếp thành công",
                "checks": {"cond": 1, "userData": 1, "mockData": 1, "confirm": 1}
            },
            {
                "utcid": "UTCID03_GetConversationStyleByPersonaId_FromSuggestedStyleId_Success",
                "type": "N",
                "cond": "Nhân vật tồn tại và có suggestedConversationStyleId hợp lệ",
                "userData": "id = 1L, suggestedConversationStyleId = 5L",
                "mockData": "personaRepositoryPort.findById(1L) -> Optional.of(persona), conversationStyleRepositoryPort.findById(5L) -> Optional.of(style)",
                "confirm": "Lấy phong cách hội thoại theo suggestedConversationStyleId thành công",
                "checks": {"cond": 2, "userData": 2, "mockData": 2, "confirm": 2}
            },
            {
                "utcid": "UTCID04_GetConversationStyleByPersonaId_NoStyleConfigured",
                "type": "N",
                "cond": "Nhân vật tồn tại nhưng không có phong cách hội thoại nào được cấu hình",
                "userData": "id = 1L, suggestedConversationStyleId = null",
                "mockData": "personaRepositoryPort.findById(1L) -> Optional.of(personaWithoutStyle)",
                "confirm": "Nhân vật không có cấu hình phong cách, trả về Optional rỗng",
                "checks": {"cond": 3, "userData": 3, "mockData": 3, "confirm": 3}
            }
        ],
        "conditions": [
            "Nhân vật không tồn tại trong cơ sở dữ liệu",
            "Nhân vật tồn tại và đối tượng phong cách hội thoại liên kết trực tiếp",
            "Nhân vật tồn tại và có suggestedConversationStyleId hợp lệ",
            "Nhân vật tồn tại nhưng không có phong cách hội thoại nào được cấu hình"
        ],
        "userDataList": [
            "id = 999L",
            "id = 1L",
            "id = 1L, suggestedConversationStyleId = 5L",
            "id = 1L, suggestedConversationStyleId = null, conversationStyle = null"
        ],
        "mockDataList": [
            "personaRepositoryPort.findById(999L) -> Optional.empty()",
            "personaRepositoryPort.findById(1L) -> Optional.of(personaWithStyle)",
            "personaRepositoryPort.findById(1L) -> Optional.of(persona), conversationStyleRepositoryPort.findById(5L) -> Optional.of(style)",
            "personaRepositoryPort.findById(1L) -> Optional.of(personaWithoutStyle)"
        ],
        "confirmList": [
            "Không tìm thấy nhân vật, trả về Optional rỗng",
            "Lấy phong cách hội thoại từ đối tượng liên kết trực tiếp thành công",
            "Lấy phong cách hội thoại theo suggestedConversationStyleId thành công",
            "Nhân vật không có cấu hình phong cách, trả về Optional rỗng"
        ]
    },
    {
        "no": 71,
        "module": "Persona",
        "method": "updatePersona",
        "class_name": "UpdatePersonaTest",
        "file_name": "UpdatePersonaTest.java",
        "requirement": "Cập nhật thông tin nhân vật và phong cách hội thoại.",
        "description": "Cập nhật thông tin tên, prompt, avatar và phong cách hội thoại của nhân vật.",
        "precondition_summary": "1. Nhân vật với ID cần cập nhật không tồn tại trong cơ sở dữ liệu\n2. Nhân vật tồn tại, chỉ cập nhật thông tin cơ bản và giữ nguyên phong cách cũ\n3. Nhân vật tồn tại, cập nhật suggestedConversationStyleId mới\n4. Nhân vật tồn tại, cập nhật thông tin phong cách hội thoại qua command",
        "cases": [
            {
                "utcid": "UTCID01_UpdatePersona_PersonaNotFound",
                "type": "A",
                "cond": "Nhân vật với ID cần cập nhật không tồn tại trong cơ sở dữ liệu",
                "userData": "UpdatePersonaCommand(id=999L, name='Updated Name', ...)",
                "mockData": "personaRepositoryPort.findById(999L) -> Optional.empty()",
                "confirm": "Ném ApplicationException mã PERSONA_NOT_FOUND (404, message: persona.not_found)",
                "checks": {"cond": 0, "userData": 0, "mockData": 0, "confirm": 0}
            },
            {
                "utcid": "UTCID02_UpdatePersona_BasicInfoOnly_Success",
                "type": "N",
                "cond": "Nhân vật tồn tại, chỉ cập nhật thông tin cơ bản và giữ nguyên phong cách cũ",
                "userData": "UpdatePersonaCommand(id=1L, name='Updated Tanaka', prompt='Updated prompt', avatarFileId=15L)",
                "mockData": "personaRepositoryPort.findById(1L) -> Optional.of(existing), personaRepositoryPort.save(Persona) -> updated",
                "confirm": "Cập nhật thông tin nhân vật thành công và giữ nguyên phong cách cũ",
                "checks": {"cond": 1, "userData": 1, "mockData": 1, "confirm": 1}
            },
            {
                "utcid": "UTCID03_UpdatePersona_WithNewStyleId_Success",
                "type": "N",
                "cond": "Nhân vật tồn tại, cập nhật suggestedConversationStyleId mới",
                "userData": "UpdatePersonaCommand(id=1L, name='Tanaka', suggestedConversationStyleId=20L)",
                "mockData": "personaRepositoryPort.findById(1L) -> Optional.of(existing), personaRepositoryPort.save(Persona) -> updated",
                "confirm": "Cập nhật thông tin nhân vật và mã phong cách hội thoại mới thành công",
                "checks": {"cond": 2, "userData": 2, "mockData": 1, "confirm": 2}
            },
            {
                "utcid": "UTCID04_UpdatePersona_WithConversationStyleCommand_Success",
                "type": "N",
                "cond": "Nhân vật tồn tại, cập nhật thông tin phong cách hội thoại qua command",
                "userData": "UpdatePersonaCommand(id=1L, name='Tanaka Sensei', conversationStyleCommand=UpdateConversationStyleCommand(...))",
                "mockData": "personaRepositoryPort.findById(1L) -> existing, conversationStyleRepositoryPort.save(Style) -> style, personaRepositoryPort.save(Persona) -> updated",
                "confirm": "Cập nhật nhân vật kèm lưu thông tin phong cách hội thoại thành công",
                "checks": {"cond": 3, "userData": 3, "mockData": 2, "confirm": 3}
            }
        ],
        "conditions": [
            "Nhân vật với ID cần cập nhật không tồn tại trong cơ sở dữ liệu",
            "Nhân vật tồn tại, chỉ cập nhật thông tin cơ bản và giữ nguyên phong cách cũ",
            "Nhân vật tồn tại, cập nhật suggestedConversationStyleId mới",
            "Nhân vật tồn tại, cập nhật thông tin phong cách hội thoại qua command"
        ],
        "userDataList": [
            "UpdatePersonaCommand(id=999L, name='Updated Name', prompt='Updated Prompt', avatarFileId=10L, suggestedConversationStyleId=1L, conversationStyleCommand=null)",
            "UpdatePersonaCommand(id=1L, name='Updated Tanaka', prompt='Updated prompt', avatarFileId=15L, suggestedConversationStyleId=null, conversationStyleCommand=null)",
            "UpdatePersonaCommand(id=1L, name='Tanaka', prompt='Prompt', avatarFileId=15L, suggestedConversationStyleId=20L, conversationStyleCommand=null)",
            "UpdatePersonaCommand(id=1L, name='Tanaka Sensei', prompt='Sensei prompt', avatarFileId=10L, conversationStyleCommand=UpdateConversationStyleCommand(...))"
        ],
        "mockDataList": [
            "personaRepositoryPort.findById(999L) -> Optional.empty()",
            "personaRepositoryPort.findById(1L) -> Optional.of(existingPersona), personaRepositoryPort.save(any(Persona.class)) -> updatedPersona",
            "conversationStyleRepositoryPort.save(any(ConversationStyle.class)) -> savedStyle, personaRepositoryPort.save(any(Persona.class)) -> updatedPersona"
        ],
        "confirmList": [
            "Ném ApplicationException mã PERSONA_NOT_FOUND (404, message: persona.not_found)",
            "Cập nhật thông tin nhân vật thành công và giữ nguyên phong cách cũ",
            "Cập nhật thông tin nhân vật và mã phong cách hội thoại mới thành công",
            "Cập nhật nhân vật kèm lưu thông tin phong cách hội thoại thành công"
        ]
    }
]

def to_roman(n):
    val = [1000, 900, 500, 400, 100, 90, 50, 40, 10, 9, 5, 4, 1]
    syb = ["m", "cm", "d", "cd", "c", "xc", "l", "xl", "x", "ix", "v", "iv", "i"]
    res = ""
    for i in range(len(val)):
        while n >= val[i]:
            res += syb[i]
            n -= val[i]
    return res

def create_excel_sheet(wb, mdef):
    sample = wb['sample']
    sheet_name = mdef['method']
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.copy_worksheet(sample)
    ws.title = sheet_name

    # Set headers
    ws['C1'] = mdef['module']
    ws['L1'] = mdef['method']
    ws['C2'] = 'TrucNV'
    ws['L2'] = 'TrucNV'
    ws['C3'] = mdef['requirement']

    num_cases = len(mdef['cases'])
    start_col = 6  # Col F
    end_col = start_col + num_cases - 1
    start_col_letter = get_column_letter(start_col)
    end_col_letter = get_column_letter(end_col)

    # Clear old data from row 7 downwards
    for r in range(7, 50):
        for c in range(1, 30):
            ws.cell(r, c).value = None

    # Row 7: Test Case IDs
    for idx, cdef in enumerate(mdef['cases']):
        col = start_col + idx
        ws.cell(7, col, f"UTCID{idx+1:02d}")

    curr_row = 8
    # 1. Condition
    ws.cell(curr_row, 1, 'Condition')
    ws.cell(curr_row, 2, 'Precondition')
    curr_row += 1

    cond_start_row = curr_row
    for idx, cond_text in enumerate(mdef['conditions']):
        ws.cell(curr_row, 2, cond_text)
        for c_idx, cdef in enumerate(mdef['cases']):
            if cdef['checks']['cond'] == idx:
                ws.cell(curr_row, start_col + c_idx, 'O')
        curr_row += 1

    # Blank row
    curr_row += 1

    # 2. userData
    ws.cell(curr_row, 1, 'userData')
    ws.cell(curr_row, 2, 'Command Data')
    curr_row += 1

    for idx, udata_text in enumerate(mdef['userDataList']):
        ws.cell(curr_row, 2, udata_text)
        for c_idx, cdef in enumerate(mdef['cases']):
            if cdef['checks']['userData'] == idx:
                ws.cell(curr_row, start_col + c_idx, 'O')
        curr_row += 1

    # Blank row
    curr_row += 1

    # 3. mockData
    ws.cell(curr_row, 1, 'mockData')
    ws.cell(curr_row, 2, 'Repository / Service')
    curr_row += 1

    for idx, mdata_text in enumerate(mdef['mockDataList']):
        ws.cell(curr_row, 2, mdata_text)
        for c_idx, cdef in enumerate(mdef['cases']):
            if cdef['checks']['mockData'] == idx:
                ws.cell(curr_row, start_col + c_idx, 'O')
        curr_row += 1

    # Blank row
    curr_row += 1

    # 4. Confirm
    ws.cell(curr_row, 1, 'Confirm')
    ws.cell(curr_row, 2, 'Expected Result')
    curr_row += 1

    for idx, conf_text in enumerate(mdef['confirmList']):
        ws.cell(curr_row, 2, conf_text)
        for c_idx, cdef in enumerate(mdef['cases']):
            if cdef['checks']['confirm'] == idx:
                ws.cell(curr_row, start_col + c_idx, 'O')
        curr_row += 1

    # Blank row
    curr_row += 1

    # 5. Result
    res_start_row = curr_row
    ws.cell(curr_row, 1, 'Result')
    ws.cell(curr_row, 2, 'Type')
    for c_idx, cdef in enumerate(mdef['cases']):
        ws.cell(curr_row, start_col + c_idx, cdef['type'])
    type_row = curr_row
    curr_row += 1

    ws.cell(curr_row, 2, 'Passed/Failed')
    for c_idx, cdef in enumerate(mdef['cases']):
        ws.cell(curr_row, start_col + c_idx, 'P')
    pf_row = curr_row
    curr_row += 1

    ws.cell(curr_row, 2, 'Executed Date')
    for c_idx, cdef in enumerate(mdef['cases']):
        ws.cell(curr_row, start_col + c_idx, datetime.datetime(2026, 8, 19, 0, 0))
    curr_row += 1

    ws.cell(curr_row, 2, 'Defect ID')
    curr_row += 1

    # Formulas in Row 5
    ws['A5'] = f'=COUNTIF({start_col_letter}{pf_row}:{end_col_letter}{pf_row},"P")'
    ws['C5'] = f'=COUNTIF({start_col_letter}{pf_row}:{end_col_letter}{pf_row},"F")'
    ws['F5'] = f'=COUNTIF({start_col_letter}{pf_row}:{end_col_letter}{pf_row},"U")'
    ws['L5'] = f'=COUNTIF({start_col_letter}{type_row}:{end_col_letter}{type_row},"N")'
    ws['M5'] = f'=COUNTIF({start_col_letter}{type_row}:{end_col_letter}{type_row},"A")'
    ws['N5'] = f'=COUNTIF({start_col_letter}{type_row}:{end_col_letter}{type_row},"B")'
    ws['O5'] = f'=COUNTA({start_col_letter}7:{end_col_letter}7)'

    # Add Data Validations
    dv_o = DataValidation(type="list", formula1='"O"', allow_blank=True)
    ws.add_data_validation(dv_o)
    dv_o.add(f'{start_col_letter}8:{end_col_letter}{type_row - 2}')

    dv_type = DataValidation(type="list", formula1='"N,A,B"', allow_blank=True)
    ws.add_data_validation(dv_type)
    dv_type.add(f'{start_col_letter}{type_row}:{end_col_letter}{type_row}')

    dv_pf = DataValidation(type="list", formula1='"P,F,U"', allow_blank=True)
    ws.add_data_validation(dv_pf)
    dv_pf.add(f'{start_col_letter}{pf_row}:{end_col_letter}{pf_row}')

    print(f'Created sheet {sheet_name} with {num_cases} test cases.')

def update_statistics(wb):
    ws = wb['Statistics']
    # Update issue date (cell F6)
    ws['F6'] = datetime.datetime(2026, 8, 19, 0, 0)

    for idx, mdef in enumerate(METHOD_DEFS):
        r = 77 + idx
        s_name = mdef['method']
        ws.cell(r, 1, mdef['no'])
        ws.cell(r, 2, f'=HYPERLINK("#{s_name}!A1", MethodList!C{r})')
        ws.cell(r, 3, f'={s_name}!A5')
        ws.cell(r, 4, f'={s_name}!C5')
        ws.cell(r, 5, f'={s_name}!F5')
        ws.cell(r, 6, f'={s_name}!L5')
        ws.cell(r, 7, f'={s_name}!M5')
        ws.cell(r, 8, f'={s_name}!N5')
        ws.cell(r, 9, f'=SUM(C{r}:H{r})')
    print('Updated Statistics sheet.')

def update_method_list(wb):
    ws = wb['MethodList']
    for idx, mdef in enumerate(METHOD_DEFS):
        r = 77 + idx
        s_name = mdef['method']
        ws.cell(r, 1, mdef['no'])
        ws.cell(r, 2, mdef['module'])
        ws.cell(r, 3, mdef['method'])
        ws.cell(r, 4, f'=HYPERLINK("#{s_name}!A1", "{s_name}")')
        ws.cell(r, 5, mdef['description'])
        ws.cell(r, 6, mdef['precondition_summary'])
    print('Updated MethodList sheet.')

def update_cover(wb):
    ws = wb['Cover']
    ws['F5'] = datetime.datetime(2026, 8, 19, 0, 0)
    ws['F6'] = datetime.datetime(2026, 8, 19, 0, 0)

    methods_str = ', '.join([m['method'] for m in METHOD_DEFS])
    r = 81
    ws.cell(r, 1, datetime.datetime(2026, 8, 19, 0, 0))
    ws.cell(r, 2, '1.0')
    ws.cell(r, 3, f'Persona ({methods_str})')
    ws.cell(r, 4, 'A')
    ws.cell(r, 5, 'Add Unit Test for methods in Persona module')
    ws.cell(r, 6, 'Persona')
    print('Updated Cover sheet.')

def generate_code_images():
    formatter = ImageFormatter(
        style='perldoc',
        font_name='Menlo',
        font_size=15,
        line_pad=6
    )

    img_map = {}
    for mdef in METHOD_DEFS:
        fpath = os.path.join(TEST_DIR, mdef['file_name'])
        with open(fpath, 'r', encoding='utf-8') as f:
            content = f.read()

        # Find all test methods
        pattern = re.compile(r'(\s+@Test[\s\S]*?\n    \})', re.MULTILINE)
        matches = pattern.findall(content)
        for m in matches:
            # find method name
            name_match = re.search(r'void\s+([A-Za-z0-9_]+)\s*\(', m)
            if name_match:
                mname = name_match.group(1)
                # Strip leading 4 spaces indent from each line
                lines = m.split('\n')
                cleaned_lines = []
                for line in lines:
                    if line.startswith('    '):
                        cleaned_lines.append(line[4:])
                    else:
                        cleaned_lines.append(line)
                cleaned_code = '\n'.join(cleaned_lines).strip()
                img_data = pygments.highlight(cleaned_code, JavaLexer(), formatter)
                out_path = os.path.join(SCRATCH_IMG_DIR, f"{mname}.png")
                with open(out_path, 'wb') as img_f:
                    img_f.write(img_data)
                img_map[mname] = out_path
                print(f'Generated image for {mname} -> {out_path}')
    return img_map

def update_docx(img_map):
    doc = docx.Document(DOCX_PATH)

    for mdef in METHOD_DEFS:
        h2_text = f"{mdef['no']}. {mdef['class_name']}"
        doc.add_paragraph(h2_text, style='Heading 2')

        for idx, cdef in enumerate(mdef['cases']):
            roman = to_roman(idx + 1)
            utcid = cdef['utcid']
            h3_text = f"{roman}. {utcid}"
            doc.add_paragraph(h3_text, style='Heading 3')

            img_path = img_map.get(utcid)
            if img_path and os.path.exists(img_path):
                p_img = doc.add_paragraph()
                r_img = p_img.add_run()
                r_img.add_picture(img_path, width=Inches(6.2))
            else:
                print(f'WARNING: Image for {utcid} not found!')

    doc.save(DOCX_PATH)
    print('Updated docx successfully.')

def main():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    for mdef in METHOD_DEFS:
        create_excel_sheet(wb, mdef)
    update_statistics(wb)
    update_method_list(wb)
    update_cover(wb)
    wb.save(EXCEL_PATH)
    print('Saved Excel file.')

    img_map = generate_code_images()
    update_docx(img_map)

if __name__ == '__main__':
    main()
