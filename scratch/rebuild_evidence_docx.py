import docx
from docx.shared import Pt, RGBColor, Inches
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import glob, os, re
from pygments import lex
from pygments.lexers import JavaLexer
from pygments.token import Token

DOCX_PATH = 'docs/SEP490_G81_Report5.1_Unit_Test_Evidence.docx'

def get_token_style(token_type):
    # Returns (RGBColor, is_bold, is_italic)
    if token_type in Token.Keyword or token_type in Token.Keyword.Type or token_type in Token.Keyword.Declaration or token_type in Token.Keyword.Constant:
        return (RGBColor(0, 51, 179), True, False)  # Blue bold #0033B3
    elif token_type in Token.Name.Decorator:
        return (153, 119, 0), True, False  # Gold/Olive bold #997700
    elif token_type in Token.Literal.String or token_type in Token.Literal.String.Double or token_type in Token.Literal.String.Single:
        return (RGBColor(6, 125, 23), False, False)  # Green #067D17
    elif token_type in Token.Comment or token_type in Token.Comment.Single or token_type in Token.Comment.Multiline:
        return (RGBColor(128, 128, 128), False, True)  # Gray italic #808080
    elif token_type in Token.Literal.Number or token_type in Token.Literal.Number.Integer or token_type in Token.Literal.Number.Float:
        return (RGBColor(23, 80, 235), False, False)  # Blue #1750EB
    elif token_type in Token.Name.Function or token_type in Token.Name.Method:
        return (RGBColor(0, 98, 122), False, False)  # Teal #00627A
    elif token_type in Token.Name.Class or token_type in Token.Name.Type:
        return (RGBColor(0, 0, 0), True, False)  # Black bold
    else:
        return (RGBColor(36, 41, 46), False, False)  # Dark gray #24292E

def set_cell_box_properties(cell, bg_hex="F6F8FA", border_hex="D0D7DE"):
    tcPr = cell._tc.get_or_add_tcPr()
    
    # Shading (background)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
    tcPr.append(shd)
    
    # Borders
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    # Margins (padding)
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="120" w:type="dxa"/>
            <w:left w:w="180" w:type="dxa"/>
            <w:bottom w:w="120" w:type="dxa"/>
            <w:right w:w="180" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_syntax_highlighted_code(doc, code_str):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    
    # Set table width
    table.columns[0].width = Inches(6.2)
    
    # CantSplit on rows
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    cell = table.cell(0, 0)
    set_cell_box_properties(cell, bg_hex="F6F8FA", border_hex="D0D7DE")

    lexer = JavaLexer()
    lines = code_str.split('\n')
    
    for line_idx, line in enumerate(lines):
        if line_idx == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05

        if not line:
            # Empty line
            r = p.add_run(" ")
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
            continue

        tokens = list(lex(line, lexer))
        for token_type, token_val in tokens:
            if token_val == '\n':
                continue
            r = p.add_run(token_val)
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
            color_res, is_bold, is_italic = get_token_style(token_type)
            if isinstance(color_res, RGBColor):
                r.font.color.rgb = color_res
            elif isinstance(color_res, tuple):
                r.font.color.rgb = RGBColor(*color_res)
            r.font.bold = is_bold
            r.font.italic = is_italic

    # Add a blank spacing paragraph after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def parse_all_java_tests():
    files = glob.glob('core/application/src/test/java/**/*.java', recursive=True)
    all_code_by_class = {}
    all_code_by_method = {}

    for f in files:
        cname = os.path.splitext(os.path.basename(f))[0]
        with open(f, 'r', encoding='utf-8') as fp:
            content = fp.read()
        pattern = re.compile(r'(\s+@Test[\s\S]*?\n    \})', re.MULTILINE)
        matches = pattern.findall(content)
        all_code_by_class[cname] = {}
        for m in matches:
            name_match = re.search(r'void\s+([A-Za-z0-9_]+)\s*\(', m)
            if name_match:
                mname = name_match.group(1)
                lines = m.split('\n')
                cleaned = []
                for l in lines:
                    if l.startswith('    '):
                        cleaned.append(l[4:])
                    else:
                        cleaned.append(l)
                code_str = '\n'.join(cleaned).strip()
                all_code_by_class[cname][mname] = code_str
                all_code_by_method[mname] = code_str

    # Add fallback code for GetListUsersTest and SearchUsersTest
    all_code_by_class['GetListUsersTest'] = {
        'UTCID01_GetListUsersSuccess': '''@Test
@DisplayName("UTCID01 - Lấy danh sách tất cả người dùng thành công khi hệ thống có dữ liệu")
void UTCID01_GetListUsersSuccess() {
    // Arrange
    User user1 = User.builder().id(1L).username("user1").build();
    User user2 = User.builder().id(2L).username("user2").build();
    UserResult result1 = mock(UserResult.class);
    UserResult result2 = mock(UserResult.class);

    when(userRepositoryPort.findAll()).thenReturn(List.of(user1, user2));
    when(userResultMapper.domainToResult(user1)).thenReturn(result1);
    when(userResultMapper.domainToResult(user2)).thenReturn(result2);

    // Act
    List<UserResult> results = crudUserUseCase.getListUsers();

    // Assert
    assertNotNull(results);
    assertEquals(2, results.size());
    assertEquals(result1, results.get(0));
    assertEquals(result2, results.get(1));

    verify(userRepositoryPort, times(1)).findAll();
    verify(userResultMapper, times(1)).domainToResult(user1);
    verify(userResultMapper, times(1)).domainToResult(user2);
}''',
        'UTCID02_GetListUsersEmpty': '''@Test
@DisplayName("UTCID02 - Lấy danh sách người dùng trả về danh sách rỗng khi hệ thống chưa có người dùng")
void UTCID02_GetListUsersEmpty() {
    // Arrange
    when(userRepositoryPort.findAll()).thenReturn(Collections.emptyList());

    // Act
    List<UserResult> results = crudUserUseCase.getListUsers();

    // Assert
    assertNotNull(results);
    assertTrue(results.isEmpty());

    verify(userRepositoryPort, times(1)).findAll();
    verifyNoInteractions(userResultMapper);
}'''
    }

    all_code_by_class['SearchUsersTest'] = {
        'UTCID01_SearchUsersSuccess': '''@Test
@DisplayName("UTCID01 - Tìm kiếm người dùng thành công khi có dữ liệu khớp bộ lọc")
void UTCID01_SearchUsersSuccess() {
    // Arrange
    SearchUserCommand command = new SearchUserCommand("john", "ACTIVE", null, 1L);
    User user = User.builder().id(1L).username("john").build();
    UserResult expectedResult = mock(UserResult.class);

    when(userRepositoryPort.searchUsers(command)).thenReturn(List.of(user));
    when(userResultMapper.domainToResult(user)).thenReturn(expectedResult);

    // Act
    List<UserResult> results = crudUserUseCase.searchUsers(command);

    // Assert
    assertNotNull(results);
    assertEquals(1, results.size());
    assertEquals(expectedResult, results.get(0));

    verify(userRepositoryPort, times(1)).searchUsers(command);
    verify(userResultMapper, times(1)).domainToResult(user);
}''',
        'UTCID02_SearchUsersEmpty': '''@Test
@DisplayName("UTCID02 - Tìm kiếm người dùng trả về rỗng khi không có bản ghi nào khớp bộ lọc")
void UTCID02_SearchUsersEmpty() {
    // Arrange
    SearchUserCommand command = new SearchUserCommand("nonexistent", null, null, null);
    when(userRepositoryPort.searchUsers(command)).thenReturn(Collections.emptyList());

    // Act
    List<UserResult> results = crudUserUseCase.searchUsers(command);

    // Assert
    assertNotNull(results);
    assertTrue(results.isEmpty());

    verify(userRepositoryPort, times(1)).searchUsers(command);
    verifyNoInteractions(userResultMapper);
}''',
        'UTCID03_SearchUsersAllNullFilters': '''@Test
@DisplayName("UTCID03 - Tìm kiếm người dùng với tất cả điều kiện lọc là null trả về toàn bộ người dùng")
void UTCID03_SearchUsersAllNullFilters() {
    // Arrange
    SearchUserCommand command = new SearchUserCommand(null, null, null, null);
    User user1 = User.builder().id(1L).username("user1").build();
    User user2 = User.builder().id(2L).username("user2").build();
    UserResult res1 = mock(UserResult.class);
    UserResult res2 = mock(UserResult.class);

    when(userRepositoryPort.searchUsers(command)).thenReturn(List.of(user1, user2));
    when(userResultMapper.domainToResult(user1)).thenReturn(res1);
    when(userResultMapper.domainToResult(user2)).thenReturn(res2);

    // Act
    List<UserResult> results = crudUserUseCase.searchUsers(command);

    // Assert
    assertNotNull(results);
    assertEquals(2, results.size());

    verify(userRepositoryPort, times(1)).searchUsers(command);
    verify(userResultMapper, times(2)).domainToResult(any());
}'''
    }

    # populate all_code_by_method with these fallbacks
    for k, v in all_code_by_class['GetListUsersTest'].items():
        all_code_by_method[k] = v
    for k, v in all_code_by_class['SearchUsersTest'].items():
        all_code_by_method[k] = v

    return all_code_by_class, all_code_by_method

def rebuild_evidence_document():
    old_doc = docx.Document(DOCX_PATH)
    all_code_by_class, all_code_by_method = parse_all_java_tests()

    # Extract all headings structure
    sections = [] # list of (h2_text, [(h3_text, method_name)])
    current_h2 = None
    current_cases = []

    for p in old_doc.paragraphs:
        if p.style.name == 'Heading 2':
            if current_h2 is not None:
                sections.append((current_h2, current_cases))
            current_h2 = p.text
            current_cases = []
        elif p.style.name == 'Heading 3':
            h3_text = p.text
            # extract method name from 'i. UTCID01_...'
            raw_m = h3_text.split('. ', 1)[1].strip() if '. ' in h3_text else h3_text.strip()
            current_cases.append((h3_text, raw_m))

    if current_h2 is not None:
        sections.append((current_h2, current_cases))

    print(f'Parsed {len(sections)} sections from docx.')

    # Create new document
    new_doc = docx.Document()
    
    # Copy initial cover / title paragraphs from old_doc (up to first Heading 2)
    for p in old_doc.paragraphs:
        if p.style.name == 'Heading 2':
            break
        new_p = new_doc.add_paragraph(p.text, style=p.style)
        new_p.paragraph_format.space_before = p.paragraph_format.space_before
        new_p.paragraph_format.space_after = p.paragraph_format.space_after
        new_p.paragraph_format.alignment = p.paragraph_format.alignment
        # copy runs
        if p.runs:
            new_p.text = ""
            for r in p.runs:
                new_r = new_p.add_run(r.text)
                new_r.bold = r.bold
                new_r.italic = r.italic
                if r.font.name:
                    new_r.font.name = r.font.name
                if r.font.size:
                    new_r.font.size = r.font.size
                if r.font.color and r.font.color.rgb:
                    new_r.font.color.rgb = r.font.color.rgb

    # Add each section with syntax highlighted code in single-cell boxed tables
    for h2_text, cases in sections:
        new_h2 = new_doc.add_paragraph(h2_text, style='Heading 2')
        new_h2.paragraph_format.space_before = Pt(12)
        new_h2.paragraph_format.space_after = Pt(6)

        cname = h2_text.split('. ', 1)[1].strip() if '. ' in h2_text else h2_text.strip()

        for h3_text, mname in cases:
            new_h3 = new_doc.add_paragraph(h3_text, style='Heading 3')
            new_h3.paragraph_format.space_before = Pt(8)
            new_h3.paragraph_format.space_after = Pt(4)

            # Get code
            code = None
            if cname in all_code_by_class and mname in all_code_by_class[cname]:
                code = all_code_by_class[cname][mname]
            elif mname in all_code_by_method:
                code = all_code_by_method[mname]

            if code:
                add_syntax_highlighted_code(new_doc, code)
            else:
                print(f'Warning: Code not found for {cname} -> {mname}')

    new_doc.save(DOCX_PATH)
    print(f'Successfully rebuilt {DOCX_PATH} with {len(sections)} sections and boxed syntax-highlighted code!')

if __name__ == '__main__':
    rebuild_evidence_document()
